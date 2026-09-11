#!/usr/bin/env python3
"""
build_resume_docx.py

Turns a structured resume-content JSON file into an ATS-friendly .docx,
and (optionally) converts it straight to PDF.

Usage:
    python3 build_resume_docx.py --content resume_content.json --out resume.docx
    python3 build_resume_docx.py --content resume_content.json --out resume.docx --pdf
    python3 build_resume_docx.py --content resume_content.json --pdf
        # --out omitted: defaults to {First}_{Last}_Resume.docx in the
        # current directory (e.g. Jane_Doe_Resume.docx), which is the
        # filename to submit to ATS portals / recruiters.
    python3 build_resume_docx.py --content resume_content.json --out resume.docx --pdf --md
        # --md also writes resume.md: a version-control-friendly source
        # mirroring the docx content (never submitted to ATS portals).

See references/data_schemas.md (in the parent skill) for the exact shape
of resume_content.json. See references/docx_creation.md for *why* each
formatting choice below exists -- it's all in service of ATS parseability
plus a clean single-page read for a human.
"""

import argparse
import json
import os
import re
import subprocess
import sys
import tempfile
import uuid

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_TAB_ALIGNMENT, WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE

BODY_FONT = "Calibri"          # Arial / Calibri / Garamond are the only ATS-safe choices
BODY_SIZE = Pt(10.5)           # never go below 10pt
NAME_SIZE = Pt(18)
HEADLINE_SIZE = Pt(11.5)
SECTION_HEAD_SIZE = Pt(11.5)
MARGIN = Inches(0.5)           # narrow margins instead of header/footer regions
USABLE_WIDTH = Inches(8.5 - 1.0)  # Letter width minus the two 0.5" margins


def set_narrow_margins(section):
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN


def set_default_font(document, name=BODY_FONT, size=BODY_SIZE):
    style = document.styles["Normal"]
    style.font.name = name
    style.font.size = size
    # Force the east-asian font slot too, otherwise some renderers fall back
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def add_bottom_border(paragraph, size=6, color="808080"):
    """Add a thin horizontal rule under a paragraph -- used for section
    headings instead of a table or a drawn line, so it stays 100% plain
    text underneath (a table row is a common ATS-parsing trap)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def tight(paragraph, before=0, after=4):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0


def add_section_heading(document, text):
    p = document.add_paragraph()
    tight(p, before=8, after=2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = SECTION_HEAD_SIZE
    add_bottom_border(p)
    return p


def add_bullet(document, text, indent=Inches(0.18), hang=Inches(0.18)):
    """A literal bullet character on a plain paragraph, not Word's
    auto-numbering. Some lightweight ATS parsers drop text that lives
    inside a numbering field -- a literal '-' + space is always read as
    plain text, so it's the safer default for a document you can't test
    against every ATS in the world."""
    p = document.add_paragraph()
    tight(p, before=0, after=2)
    pf = p.paragraph_format
    pf.left_indent = indent
    pf.first_line_indent = -hang
    run = p.add_run(f"-  {text}")
    run.font.size = BODY_SIZE
    return p


def add_right_tab_line(document, left_text, right_text, bold_left=True):
    """One line with left-aligned text and a right-aligned date.

    The visual alignment uses a right tab stop, but ATS text extraction
    often drops tab characters entirely. To prevent the title and date
    from merging into one token (e.g. "Engineer07/2025"), the date run
    is prefixed with an explicit " | " separator plus the tab — so even
    if the tab is stripped, a pipe + spaces remain between the tokens.
    """
    p = document.add_paragraph()
    tight(p, before=6, after=0)
    p.paragraph_format.tab_stops.add_tab_stop(USABLE_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run(left_text)
    r1.bold = bold_left
    r1.font.size = BODY_SIZE
    if right_text and right_text.strip():
        # Pipe is the fallback separator if the tab is lost in parsing;
        # tab provides the visual right-alignment for human readers.
        r2 = p.add_run(f"  |  \t{right_text.strip()}")
        r2.bold = False
        r2.font.size = BODY_SIZE
    return p


def add_hyperlink(paragraph, url, text, color="0B5394", underline=True):
    """Insert a native clickable hyperlink into a python-docx paragraph.
    Allows human recruiters in 2026 to click directly through to LinkedIn,
    GitHub, or live project demos without breaking ATS parsers."""
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rPr.append(c)

    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)

    new_run.append(rPr)
    new_run.text = text
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_contact_line(document, contact):
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, before=0, after=8)

    entries = []
    if contact.get("location"):
        entries.append(("text", contact["location"]))
    if contact.get("phone"):
        entries.append(("text", contact["phone"]))
    if contact.get("email"):
        email = contact["email"]
        target = f"mailto:{email}" if not email.startswith("mailto:") else email
        entries.append(("link", email, target))
    for key in ("linkedin", "github", "website"):
        if contact.get(key):
            raw = contact[key].strip()
            target = raw if raw.startswith(("http://", "https://")) else f"https://{raw}"
            entries.append(("link", raw, target))

    for idx, item in enumerate(entries):
        if idx > 0:
            sep = p.add_run("  |  ")
            sep.font.size = Pt(10)
        if item[0] == "text":
            run = p.add_run(item[1])
            run.font.size = Pt(10)
        elif item[0] == "link":
            add_hyperlink(p, item[2], item[1], color="0B5394", underline=True)


def build_header(document, content):
    name_p = document.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(name_p, before=0, after=0)
    name_run = name_p.add_run(content["contact"]["name"])
    name_run.bold = True
    name_run.font.size = NAME_SIZE

    add_contact_line(document, content["contact"])


def build_summary(document, content):
    # The headline itself IS the section heading here (a slightly more
    # descriptive title than a bare "Professional Summary" label), with
    # the actual summary sentence(s) directly underneath.
    add_section_heading(document, content["headline"])
    p = document.add_paragraph()
    tight(p, before=0, after=6)
    run = p.add_run(content["summary"])
    run.font.size = BODY_SIZE


def build_skills(document, content):
    add_section_heading(document, "Skills")
    for row in content.get("skills", []):
        p = document.add_paragraph()
        tight(p, before=0, after=2)
        label = p.add_run(f"{row['category']}: ")
        label.bold = True
        label.font.size = BODY_SIZE
        items = p.add_run(" | ".join(row["items"]))
        items.font.size = BODY_SIZE


def build_experience(document, content):
    jobs = content.get("experience", [])
    if not jobs:
        return
    add_section_heading(document, "Work Experience")
    for job in jobs:
        left = f"{job['company']}, {job['location']} | {job['title']}"
        right = f"{job['start']} - {job['end']}"
        add_right_tab_line(document, left, right)
        for bullet in job.get("bullets", []):
            add_bullet(document, bullet)


def build_education(document, content):
    add_section_heading(document, "Education")
    for edu in content.get("education", []):
        left = f"{edu['degree']}, {edu['school']}, {edu['location']}"
        right = edu.get("grad_year", "")
        add_right_tab_line(document, left, right)
        details = []
        if edu.get("gpa"):
            details.append(f"GPA: {edu['gpa']}")
        details.extend(edu.get("achievements", []))
        if details:
            p = document.add_paragraph()
            tight(p, before=0, after=4)
            run = p.add_run(" | ".join(details))
            run.font.size = BODY_SIZE


def split_description_to_bullets(description, max_bullets=3):
    """Fallback: turn a prose project description into achievement bullets.

    Projects should arrive with explicit 2-4 "bullets" shaped like work
    experience bullets (see resume_writing_rules.md section 8). Older
    profiles only have a single "description" paragraph — rendering that
    as prose breaks the section's visual rhythm and parses as one blob.
    Splitting on sentence / clause boundaries keeps the same bulleted
    style even for legacy data, without inventing new content.
    """
    text = (description or "").strip()
    if not text:
        return []
    parts = re.split(r"(?<=[.!?])\s+|;\s+|\s\|\s+", text)
    bullets = [p.strip().rstrip(" |;").strip() for p in parts]
    bullets = [b for b in bullets if b]
    if len(bullets) > max_bullets:
        bullets = bullets[: max_bullets - 1] + [" ".join(bullets[max_bullets - 1 :])]
    return bullets


def build_projects(document, content):
    projects = content.get("projects", [])
    if not projects:
        return
    add_section_heading(document, "Projects")
    for proj in projects:
        p = document.add_paragraph()
        tight(p, before=4, after=0)
        name_run = p.add_run(proj["name"])
        name_run.bold = True
        name_run.font.size = BODY_SIZE
        if proj.get("url"):
            raw_url = proj["url"].strip()
            target_url = raw_url if raw_url.startswith(("http://", "https://")) else f"https://{raw_url}"
            open_run = p.add_run("  (")
            open_run.font.size = Pt(10)
            add_hyperlink(p, target_url, raw_url, color="0B5394", underline=True)
            close_run = p.add_run(")")
            close_run.font.size = Pt(10)
        bullets = list(proj.get("bullets") or [])
        if not bullets and proj.get("description"):
            bullets = split_description_to_bullets(proj["description"])
            print(
                f"Warning: project '{proj.get('name', '?')}' has no 'bullets'; "
                f"split description into {len(bullets)} bullet(s). "
                f"Prefer explicit 2-4 achievement bullets per data_schemas.md.",
                file=sys.stderr,
            )
        for bullet in bullets:
            add_bullet(document, bullet)
        if bullets:
            # breathing room after the last bullet before the next project
            document.paragraphs[-1].paragraph_format.space_after = Pt(4)


def build_awards(document, content):
    awards = content.get("awards", [])
    if not awards:
        return
    add_section_heading(document, "Awards, Accolades and Certifications")
    for award in awards:
        p = document.add_paragraph()
        tight(p, before=0, after=2)
        run = p.add_run(f"{award['year']} | {award['achievement']} | {award['name']}")
        run.font.size = BODY_SIZE


SECTION_BUILDERS = {
    "contact": build_header,
    "headline_summary": build_summary,
    "skills": build_skills,
    "experience": build_experience,
    "education": build_education,
    "projects": build_projects,
    "awards": build_awards,
}

DEFAULT_ORDER = [
    "contact",
    "headline_summary",
    "skills",
    "experience",
    "education",
    "projects",
    "awards",
]


def default_resume_stem(contact_name):
    """Derive a human-friendly file stem from the candidate's full name.

    Returns e.g. "Jane_Doe_Resume" for "Jane Doe" — the filename to
    submit to ATS portals / recruiters. Falls back to "Resume" when no
    usable name is present. Only alphanumerics + underscore are kept so
    the name is safe on every OS and never confuses an ATS upload form.
    """
    tokens = re.findall(r"[A-Za-z0-9]+", (contact_name or "").strip())
    if not tokens:
        return "Resume"
    if len(tokens) == 1:
        return f"{tokens[0]}_Resume"
    return f"{tokens[0]}_{tokens[-1]}_Resume"


def build_markdown(content):
    """Version-control-friendly markdown source of the same resume.

    Mirrors the docx section order and content one-to-one so diffs show
    what actually changed between tailored versions. Not submitted to
    ATS portals — the docx/PDF remain the deliverables.
    """
    lines = []
    contact = content.get("contact", {})
    lines.append(f"# {contact.get('name', '')}")
    bits = [contact.get(k, "") for k in ("location", "phone", "email", "linkedin", "github", "website")]
    lines.append(" | ".join(b for b in bits if b))
    lines.append("")
    lines.append(f"## {content.get('headline', '')}")
    lines.append("")
    lines.append(content.get("summary", ""))
    lines.append("")
    lines.append("## Skills")
    lines.append("")
    for row in content.get("skills", []):
        lines.append(f"- **{row.get('category', '')}:** {' | '.join(row.get('items', []))}")
    lines.append("")
    lines.append("## Work Experience")
    lines.append("")
    for job in content.get("experience", []):
        lines.append(
            f"### {job.get('company', '')}, {job.get('location', '')} | "
            f"{job.get('title', '')} | {job.get('start', '')} - {job.get('end', '')}"
        )
        lines.append("")
        for bullet in job.get("bullets", []):
            lines.append(f"- {bullet}")
        lines.append("")
    lines.append("## Education")
    lines.append("")
    for edu in content.get("education", []):
        lines.append(
            f"### {edu.get('degree', '')}, {edu.get('school', '')}, {edu.get('location', '')} | "
            f"{edu.get('grad_year', '')}"
        )
        lines.append("")
        details = []
        if edu.get("gpa"):
            details.append(f"GPA: {edu['gpa']}")
        details.extend(edu.get("achievements", []))
        if details:
            lines.append(f"{' | '.join(details)}")
            lines.append("")
    projects = content.get("projects", [])
    if projects:
        lines.append("## Projects")
        lines.append("")
        for proj in projects:
            title = f"### {proj.get('name', '')}"
            if proj.get("url"):
                title += f" ({proj['url']})"
            lines.append(title)
            lines.append("")
            bullets = list(proj.get("bullets") or [])
            if not bullets and proj.get("description"):
                bullets = split_description_to_bullets(proj["description"])
            for bullet in bullets:
                lines.append(f"- {bullet}")
            lines.append("")
    awards = content.get("awards", [])
    if awards:
        lines.append("## Awards, Accolades and Certifications")
        lines.append("")
        for award in awards:
            lines.append(f"- {award.get('year', '')} | {award.get('achievement', '')} | {award.get('name', '')}")
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


def build_resume(content, out_path):
    document = Document()
    section = document.sections[0]
    set_narrow_margins(section)
    set_default_font(document)

    order = content.get("section_order", DEFAULT_ORDER)
    # contact/header always renders first regardless of requested order
    build_header(document, content)
    for key in order:
        if key in ("contact",):
            continue
        builder = SECTION_BUILDERS.get(key)
        if builder:
            builder(document, content)

    document.save(out_path)
    return out_path


def convert_to_pdf(docx_path, out_dir=None):
    out_dir = out_dir or os.path.dirname(os.path.abspath(docx_path)) or "."
    # Give this conversion its own LibreOffice user profile dir. Headless
    # soffice locks a single shared profile by default, which throws
    # spurious errors when an agent runs conversions back-to-back or in
    # parallel -- a fresh temp profile per call sidesteps that entirely.
    profile_dir = os.path.join(tempfile.gettempdir(), f"lo_profile_{uuid.uuid4().hex}")
    env = os.environ.copy()
    cmd = [
        "soffice",
        "--headless",
        "--norestore",
        f"-env:UserInstallation=file://{profile_dir}",
        "--convert-to",
        "pdf",
        "--outdir",
        out_dir,
        docx_path,
    ]
    result = subprocess.run(cmd, capture_output=True, text=True, env=env, timeout=120)
    if result.returncode != 0:
        raise RuntimeError(
            f"soffice conversion failed (code {result.returncode}):\n"
            f"stdout: {result.stdout}\nstderr: {result.stderr}"
        )
    pdf_path = os.path.join(
        out_dir, os.path.splitext(os.path.basename(docx_path))[0] + ".pdf"
    )
    if not os.path.exists(pdf_path):
        raise RuntimeError(f"Expected PDF not found at {pdf_path} after conversion")
    return pdf_path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--content", required=True, help="Path to resume_content.json")
    parser.add_argument(
        "--out",
        required=False,
        default=None,
        help="Output .docx path (default: {First}_{Last}_Resume.docx from contact.name)",
    )
    parser.add_argument("--pdf", action="store_true", help="Also convert to PDF")
    parser.add_argument("--md", action="store_true",
                        help="Also write a markdown source next to the docx")
    args = parser.parse_args()

    with open(args.content, "r", encoding="utf-8") as f:
        content = json.load(f)

    out_path = args.out
    if not out_path:
        stem = default_resume_stem(content.get("contact", {}).get("name", ""))
        out_path = f"{stem}.docx"
    parent = os.path.dirname(os.path.abspath(out_path))
    if parent:
        os.makedirs(parent, exist_ok=True)

    docx_path = build_resume(content, out_path)
    print(f"Wrote {docx_path}")

    if args.md:
        md_path = os.path.splitext(docx_path)[0] + ".md"
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(build_markdown(content))
        print(f"Wrote {md_path}")

    if args.pdf:
        pdf_path = convert_to_pdf(docx_path)
        print(f"Wrote {pdf_path}")


if __name__ == "__main__":
    main()
